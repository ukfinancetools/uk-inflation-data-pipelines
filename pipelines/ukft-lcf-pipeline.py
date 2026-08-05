#!/usr/bin/env python3
"""UKFT RCOL - LCF (Living Costs & Food Survey) expenditure-weights pipeline.
LCF/Family Spending is spreadsheet-only (annual, ~Sep release, ~12-18 month lag).
Resolves the latest Family Spending workbooks 2 (income/composition/tenure) and 3
(region) from the ONS dataset landing pages, parses high-level expenditure (13
COICOP categories) per household group, derives renormalised CPI-division weights
(categories 1-12; category 13 'other expenditure items' excluded and disclosed),
and writes ukft-lcf-weights.json for the RCOL personal inflation calculator.

Profile pre-fill is SINGLE-dimension by design: the published tables are
univariate; joint profiles are not derivable from published aggregates.
Annual cadence - run manually on ONS publication (or via the planned
RCOL non-API refresh task). Mirrors ukft-hci-pipeline.py conventions.
Usage: python3 ukft-lcf-pipeline.py [--dry-run]
"""
import sys, os, re, json, io, datetime, urllib.request
import openpyxl

UA = {'User-Agent': 'curl/8.5.0', 'Accept': '*/*'}
DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'data')
DRY = '--dry-run' in sys.argv
LANDING_WB2 = "https://www.ons.gov.uk/peoplepopulationandcommunity/personalandhouseholdfinances/expenditure/datasets/familyspendingworkbook2expenditurebyincome"
LANDING_WB3 = "https://www.ons.gov.uk/peoplepopulationandcommunity/personalandhouseholdfinances/expenditure/datasets/familyspendingworkbook3expenditurebyregion"
FALLBACK_WB2 = "https://www.ons.gov.uk/file?uri=/peoplepopulationandcommunity/personalandhouseholdfinances/expenditure/datasets/familyspendingworkbook2expenditurebyincome/fye2024/workbook2expenditurebyincome.xlsx"
FALLBACK_WB3 = "https://www.ons.gov.uk/file?uri=/peoplepopulationandcommunity/personalandhouseholdfinances/expenditure/datasets/familyspendingworkbook3expenditurebyregion/fye2024/workbook3expenditurebyregion.xlsx"

CATS = [  # key MUST match the divisions block keys in ukft-inflation-data.json
    ('d01','01','Food & non-alcoholic drinks'), ('d02','02','Alcoholic drinks & tobacco'),
    ('d03','03','Clothing & footwear'), ('d04','04','Housing, water, fuel & power'),
    ('d05','05','Furniture & household goods'), ('d06','06','Health'),
    ('d07','07','Transport'), ('d08','08','Communication'),
    ('d09','09','Recreation & culture'), ('d10','10','Education'),
    ('d11','11','Restaurants & hotels'), ('d12','12','Miscellaneous goods & services'),
]

# (sheet, expected numeric col count in a category row, [(col_index_0based_among_numeric, profile_key, label, sublabel)])
WB2_TABLES = [
    ('A4', 11, [(i, f'decile_{i+1:02d}', f'Income decile {i+1}', None) for i in range(10)] + [(10,'uk_average','UK average household',None)]),
    ('A23', 13, [
        (0,'comp_ret_sp_single','Retired single (mainly state pension)',None),
        (1,'comp_ret_sp_couple','Retired couple (mainly state pension)',None),
        (2,'comp_ret_single','Retired single',None),
        (3,'comp_ret_couple','Retired couple',None),
        (4,'comp_single','Working-age single',None),
        (5,'comp_couple','Working-age couple',None),
        (6,'comp_1a_1c','Single parent, one child',None),
        (7,'comp_1a_2c','Single parent, two or more children',None),
        (8,'comp_2a_1c','Couple, one child',None),
        (9,'comp_2a_2c','Couple, two children',None),
        (10,'comp_2a_3c','Couple, three or more children',None),
        (11,'comp_3a_noch','Three or more adults, no children',None),
        (12,'comp_3a_ch','Three or more adults, with children',None)]),
    ('A32', 11, [
        (0,'tenure_outright','Own outright',None),
        (1,'tenure_mortgage','Buying with a mortgage',None),
        (6,'tenure_social','Social renter',None),
        (9,'tenure_private','Private renter',None)]),
]
WB3_TABLES = [
    ('A33', 14, [
        (2,'region_north_east','North East',None),(3,'region_north_west','North West',None),
        (4,'region_yorkshire','Yorkshire and The Humber',None),(5,'region_east_midlands','East Midlands',None),
        (6,'region_west_midlands','West Midlands',None),(7,'region_east','East of England',None),
        (8,'region_london','London',None),(9,'region_south_east','South East',None),
        (10,'region_south_west','South West',None),(11,'region_wales','Wales',None),
        (12,'region_scotland','Scotland',None),(13,'region_northern_ireland','Northern Ireland',None)]),
]

def fetch(url):
    return urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=40).read()

def resolve_url(landing, fallback, slugbit):
    try:
        html = fetch(landing).decode('utf-8','replace')
        hrefs = re.findall(r'(/file\?uri=[^"\']*' + slugbit + r'[^"\']*\.xlsx)', html)
        hrefs = [h for h in hrefs if '/current/' not in h]
        if hrefs: return "https://www.ons.gov.uk"+hrefs[0]
    except Exception as e:
        print(f"  WARN url-resolve failed ({e}); using fallback")
    return fallback

def num(v):
    if v is None: return None
    s = str(v).strip().replace(',','')
    if s in ('..','[x]','[c]','-',':','') : return None
    try: return float(s)
    except ValueError: return None

def grid(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]

def cat_rows(g):
    """Find the 13 category rows + total row. Returns {catno:int -> row}, total_row."""
    out, total = {}, None
    for r in g:
        a = str(r[0]).strip() if r and r[0] is not None else ''
        if re.fullmatch(r'\d{1,2}', a) and 1 <= int(a) <= 13 and r[1]:
            out[int(a)] = r
        joined = ' '.join(str(c) for c in r[:5] if c is not None)
        if total is None and joined.strip().startswith('Total expenditure'):
            total = r
    if len(out) != 13 or total is None:
        raise ValueError(f"category parse failed: got {sorted(out)} total={'yes' if total else 'no'}")
    return out, total

def numeric_cols(row):
    """Column indices holding numeric values in a category row (values start after label cols)."""
    return [i for i,v in enumerate(row) if i >= 2 and num(v) is not None]

def parse_table(g, expected_cols, picks, sheet):
    cats, total_row = cat_rows(g)
    # use category 1 (Food - never suppressed) to fix the column map
    cols = numeric_cols(cats[1])
    if len(cols) != expected_cols:
        raise ValueError(f"{sheet}: expected {expected_cols} numeric cols, got {len(cols)}")
    profiles = {}
    for idx, key, label, sub in picks:
        col = cols[idx]
        weekly, est = {}, []
        for cno,(ckey,_,_) in zip(range(1,13), CATS):
            v = num(cats[cno][col])
            if v is None: est.append(ckey); v = -1.0  # placeholder; substituted later
            weekly[ckey] = v
        c13 = num(cats[13][col]); tot = num(total_row[col])
        profiles[key] = {'label': label, 'weekly': weekly, 'cat13': c13, 'total': tot, 'estimated': est}
    return profiles, cats

def boundaries(g):
    for r in g:
        lead = ' '.join(str(c) for c in r[:5] if c is not None)
        if lead.strip().startswith('Lower boundary'):
            return [num(v) for v in r if num(v) is not None]
    return []

def main():
    print("UKFT RCOL - LCF weights pipeline")
    u2 = resolve_url(LANDING_WB2, FALLBACK_WB2, 'workbook2'); print(f"  wb2: {u2}")
    u3 = resolve_url(LANDING_WB3, FALLBACK_WB3, 'workbook3'); print(f"  wb3: {u3}")
    vintage_m2 = re.search(r'/(fye\d{4})/', u2)
    vintage = vintage_m2.group(1).upper().replace('FYE','FYE ') if vintage_m2 else 'unknown'
    wb2 = openpyxl.load_workbook(io.BytesIO(fetch(u2)), read_only=True)
    wb3 = openpyxl.load_workbook(io.BytesIO(fetch(u3)), read_only=True)

    profiles = {}
    decile_bounds = []
    for sheet, n, picks in WB2_TABLES:
        g = grid(wb2[sheet]); p, cats = parse_table(g, n, picks, sheet)
        profiles.update(p)
        if sheet == 'A4': decile_bounds = boundaries(g)
        print(f"  {sheet}: {len(p)} profiles")
    for sheet, n, picks in WB3_TABLES:
        g = grid(wb3[sheet]); p, _ = parse_table(g, n, picks, sheet)
        profiles.update(p)
        print(f"  {sheet}: {len(p)} profiles (3-year-average basis)")

    # substitute suppressed values with the UK-average weekly value, then weights
    uk = profiles['uk_average']
    if uk['estimated']: raise ValueError("UK average has suppressed cells - parse wrong")
    for key, pr in profiles.items():
        for ckey in pr['estimated']:
            pr['weekly'][ckey] = round(uk['weekly'][ckey], 2)
        s = sum(pr['weekly'].values())
        if s <= 0: raise ValueError(f"{key}: non-positive total")
        w = {k: 100.0*v/s for k,v in pr['weekly'].items()}
        w = {k: round(v,1) for k,v in w.items()}
        drift = round(100.0 - sum(w.values()), 1)
        if abs(drift) >= 0.05:  # pin rounding drift onto the largest weight
            big = max(w, key=w.get); w[big] = round(w[big] + drift, 1)
        pr['weights'] = w
        pr['total_1_12'] = round(s, 2)

    # decile sublabels from gross boundaries (9 boundaries: deciles 2-10 lower bounds)
    if len(decile_bounds) >= 9:
        b = [int(x) for x in decile_bounds[:9]]
        subs = [f"under GBP{b[0]}/wk gross"] + [f"GBP{b[i]}-{b[i+1]}/wk gross" for i in range(8)] + [f"over GBP{b[8]}/wk gross"]
        for i in range(10):
            profiles[f'decile_{i+1:02d}']['sublabel'] = subs[i]

    doc = {
        '_meta': {
            'description': 'UKFT RCOL Layer 3 - LCF household expenditure weights by single-dimension profile (CPI divisions 1-12, renormalised; category 13 excluded)',
            'generated': datetime.datetime.utcnow().isoformat()+'Z',
            'vintage': vintage, 'published': '2025-09-10' if vintage=='FYE 2024' else None,
            'source': [u2, u3], 'licence': 'Open Government Licence v3.0',
            'attribution': 'Office for National Statistics, Living Costs and Food Survey',
            'caveats': [
                'Weights derived from LCF average weekly expenditure per group; categories 1-12 only - category 13 (mortgage interest, council tax, other non-consumption items) is excluded and weights renormalised.',
                'Profiles are single-dimension (published LCF tables are univariate); joint profiles are not derivable from published aggregates.',
                'Region profiles are 3-year pooled averages; other profiles are single-year.',
                'Suppressed cells (Health/Education in some groups) substituted with the UK-average value and flagged in estimated.',
                'Income decile profiles are GROSS income deciles (not equivalised disposable - a different basis from the ONS Household Costs Indices deciles).',
            ],
        },
        'categories': [{'key':k,'division':d,'label':l} for k,d,l in CATS],
        'profiles': profiles,
    }
    out = os.path.join(DATA_DIR, 'ukft-lcf-weights.json')
    if DRY: print(f"  [DRY RUN] would write {out}"); return
    with open(out,'w',encoding='utf-8') as f: json.dump(doc,f,indent=1)
    print(f"  ok ukft-lcf-weights.json ({os.path.getsize(out)/1024:.1f} KB) | {len(profiles)} profiles | vintage {vintage}")

if __name__ == '__main__':
    main()
